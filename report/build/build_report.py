"""Builds the SIWES IT project report (Word + PDF).

    python report/build/make_notebooks.py   # charts from the analysis notebooks
    python report/build/make_figures.py     # diagrams + code snippets
    python report/build/build_report.py     # this file -> report/output/*.docx, *.pdf

Two passes: the first build is converted to PDF to find the page of every heading and figure,
the second build writes those page numbers into the table of contents and list of figures.
"""
import json
import re
import subprocess
from pathlib import Path

import pdfplumber
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image

from content import CHAPTERS, PRELIMS, REFERENCES, APPENDICES

ROOT = Path(__file__).resolve().parents[1]
REPO = ROOT.parent
OUT = ROOT / "output"
TMP = ROOT / "build" / "_tmp"
OUT.mkdir(exist_ok=True); TMP.mkdir(exist_ok=True)
NAME = "Joseph_Cosmos_Aniebiet_SIWES_IT_Project_Report"
TEXT_W = 6.5   # inches between 1" margins on US Letter (same page set-up as the department template)


# ------------------------------------------------------------------ low-level helpers
def set_cell_bg(cell, hex_color):
    tc = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_color)
    tc.append(shd)


def add_field(run, instr):
    for kind, text in (("begin", None), (None, instr), ("separate", None), (None, "1"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), kind); run._r.append(el)
        elif text == instr:
            el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = instr; run._r.append(el)
        else:
            el = OxmlElement("w:t"); el.text = text; run._r.append(el)


def page_numbers(section, fmt, start):
    pg = section._sectPr.find(qn("w:pgNumType"))
    if pg is None:
        pg = OxmlElement("w:pgNumType"); section._sectPr.append(pg)
    pg.set(qn("w:fmt"), fmt); pg.set(qn("w:start"), str(start))


def keep_with_next(p):
    p.paragraph_format.keep_with_next = True


def crop_image(src, box, name):
    out = TMP / f"{name}.png"
    Image.open(src).crop(box).save(out, optimize=True)
    return out


def side_by_side(paths, name, gap=40, bg="white"):
    out = TMP / f"{name}.png"
    ims = [Image.open(p).convert("RGB") for p in paths]
    h = min(im.height for im in ims)
    ims = [im.resize((int(im.width * h / im.height), h)) for im in ims]
    W = sum(im.width for im in ims) + gap * (len(ims) - 1)
    canvas = Image.new("RGB", (W, h), bg); x = 0
    for im in ims:
        canvas.paste(im, (x, 0)); x += im.width + gap
    canvas.save(out, optimize=True)
    return out


# ------------------------------------------------------------------ report writer
class Report:
    def __init__(self, pages=None):
        self.pages = pages or {}          # key -> printed page number (from pass 1)
        self.doc = Document()
        self.toc, self.lof, self.lot = [], [], []
        self.chapter = 0; self.fig_n = 0; self.tab_n = 0
        self._styles()

    # ---------- styles
    def _styles(self):
        d = self.doc
        s = d.sections[0]
        s.page_width, s.page_height = Inches(8.5), Inches(11)
        for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(s, side, Inches(1))
        st = d.styles["Normal"]
        st.font.name = "Times New Roman"; st.font.size = Pt(12)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        pf = st.paragraph_format
        pf.line_spacing = 1.5; pf.space_after = Pt(6); pf.space_before = Pt(0)
        for name in ("List Bullet", "List Number"):
            ls = d.styles[name]; ls.font.name = "Times New Roman"; ls.font.size = Pt(12)
            ls.paragraph_format.line_spacing = 1.5; ls.paragraph_format.space_after = Pt(3)
        for lvl, size in ((1, 14), (2, 12), (3, 12)):
            h = d.styles[f"Heading {lvl}"]
            h.font.name = "Times New Roman"; h.font.size = Pt(size); h.font.bold = True
            h.font.color.rgb = RGBColor(0, 0, 0); h.font.italic = (lvl == 3)
            h.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
            rf = h.element.rPr.rFonts
            for a in ("w:ascii", "w:hAnsi", "w:cs"):
                rf.set(qn(a), "Times New Roman")
            for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
                rf.attrib.pop(qn(a), None)
            h.paragraph_format.space_before = Pt(12 if lvl > 1 else 0); h.paragraph_format.space_after = Pt(6)
            h.paragraph_format.line_spacing = 1.5; h.paragraph_format.keep_with_next = True

    # ---------- text
    def _runs(self, p, text, size=None, italic=False):
        # **bold** and `code` inline markup
        for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
            if not part:
                continue
            if part.startswith("**"):
                r = p.add_run(part[2:-2]); r.bold = True
            elif part.startswith("`"):
                r = p.add_run(part[1:-1]); r.font.name = "Courier New"; r.font.size = Pt(10.5)
            else:
                r = p.add_run(part)
            if size and not part.startswith("`"):
                r.font.size = Pt(size)
            if italic:
                r.italic = True
        return p

    def para(self, text, align="justify", size=None, italic=False, space_after=None):
        p = self.doc.add_paragraph()
        p.alignment = {"justify": WD_ALIGN_PARAGRAPH.JUSTIFY, "center": WD_ALIGN_PARAGRAPH.CENTER,
                       "left": WD_ALIGN_PARAGRAPH.LEFT, "right": WD_ALIGN_PARAGRAPH.RIGHT}[align]
        if space_after is not None:
            p.paragraph_format.space_after = Pt(space_after)
        return self._runs(p, text, size, italic)

    def bullets(self, items, numbered=False):
        for it in items:
            p = self.doc.add_paragraph(style="List Number" if numbered else "List Bullet")
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            if ":" in it and not it.startswith("**") and it.index(":") < 60:
                head, rest = it.split(":", 1)
                self._runs(p, f"**{head}:**{rest}")
            else:
                self._runs(p, it)

    def h1(self, number, title):
        self.chapter = number; self.fig_n = 0; self.tab_n = 0
        p = self.doc.add_paragraph(style="Heading 1"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.page_break_before = not self._after_section_break(p)
        p.add_run(f"CHAPTER {['', 'ONE', 'TWO', 'THREE'][number]}")
        p2 = self.doc.add_paragraph(style="Heading 1"); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.add_run(title.upper())
        p2.paragraph_format.space_after = Pt(18)
        self.toc.append((0, f"Chapter {number}: {title}", f"h1-{number}"))

    def h_plain(self, title, key, level=0):
        """Unnumbered heading (prelims, references, appendices)."""
        p = self.doc.add_paragraph(style="Heading 1"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.page_break_before = not self._after_section_break(p)
        p.add_run(title.upper()); p.paragraph_format.space_after = Pt(18)
        if key:
            self.toc.append((level, title, key))

    def _after_section_break(self, p):
        prev = p._p.getprevious()
        return prev is not None and prev.find(qn("w:pPr")) is not None and \
            prev.find(qn("w:pPr")).find(qn("w:sectPr")) is not None

    def h2(self, text):
        p = self.doc.add_paragraph(style="Heading 2"); p.add_run(text.upper() if not text[0].isdigit() else
                                                                 re.sub(r"^([\d.]+)\s+(.*)$", lambda m: f"{m.group(1)}\t{m.group(2).upper()}", text))
        p.paragraph_format.tab_stops.add_tab_stop(Inches(0.6))
        if text[0].isdigit():
            self.toc.append((1, text, "h2-" + text.split()[0]))

    def h3(self, text):
        p = self.doc.add_paragraph(style="Heading 3"); p.add_run(text)
        if text[0].isdigit():
            self.toc.append((2, text, "h3-" + text.split()[0]))

    # ---------- figures and tables
    def figure(self, path, caption, width=6.0, max_h=7.6):
        path = Path(path)
        if not path.exists():
            raise FileNotFoundError(path)
        im = Image.open(path)
        w = width
        if path.parent.name == "snippets":
            w = min(6.3, im.width * 6.0 / 1450)   # code images: same font size everywhere (~8 pt)
        elif path.parent.name == "figures" and path.name.startswith(("ac_", "bp_")):
            w = min(w, 5.7)
        if im.height / im.width * w > max_h:
            w = max_h * im.width / im.height
        p = self.doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2); p.paragraph_format.line_spacing = 1.0
        keep_with_next(p)
        p.add_run().add_picture(str(path), width=Inches(w))
        self.fig_n += 1
        label = f"Figure {self.chapter}.{self.fig_n}" if self.chapter else f"Figure {self.fig_n}"
        c = self.doc.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        c.paragraph_format.space_after = Pt(10); c.paragraph_format.line_spacing = 1.15
        r = c.add_run(f"{label}: "); r.bold = True; r.font.size = Pt(10.5)
        r2 = c.add_run(caption); r2.italic = True; r2.font.size = Pt(10.5)
        self.lof.append((label, caption, "fig-" + label.split()[1]))

    def table(self, headers, rows, caption, widths=None, size=10.5):
        self.tab_n += 1
        label = f"Table {self.chapter}.{self.tab_n}"
        c = self.doc.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        c.paragraph_format.space_after = Pt(4); keep_with_next(c)
        r = c.add_run(f"{label}: "); r.bold = True; r.font.size = Pt(10.5)
        r2 = c.add_run(caption); r2.italic = True; r2.font.size = Pt(10.5)
        self.lot.append((label, caption, "tab-" + label.split()[1]))
        t = self.doc.add_table(rows=1, cols=len(headers)); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, h in enumerate(headers):
            cell = t.rows[0].cells[i]; cell.text = ""
            run = cell.paragraphs[0].add_run(h); run.bold = True; run.font.size = Pt(size); run.font.color.rgb = RGBColor(255, 255, 255)
            set_cell_bg(cell, "0B2A6F")
        for ri, row in enumerate(rows):
            cells = t.add_row().cells
            for i, val in enumerate(row):
                cells[i].text = ""
                para = cells[i].paragraphs[0]; para.paragraph_format.line_spacing = 1.1; para.paragraph_format.space_after = Pt(0)
                self._runs(para, str(val), size)
                if ri % 2:
                    set_cell_bg(cells[i], "EEF2F8")
        for row in t.rows:
            for c_ in row.cells:
                for para in c_.paragraphs:
                    para.paragraph_format.line_spacing = 1.1; para.paragraph_format.space_after = Pt(1)
        if widths:
            t.autofit = False
            scale = TEXT_W / sum(widths) if sum(widths) > TEXT_W else 1.0
            for i, w in enumerate(widths):
                t.columns[i].width = Inches(w * scale)
                for row in t.rows:
                    row.cells[i].width = Inches(w * scale)
        sp = self.doc.add_paragraph(); sp.paragraph_format.space_after = Pt(4); sp.paragraph_format.line_spacing = 1.0

    def page_break(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ---------- body blocks from content.py
    def number(self, blocks):
        """Pre-pass: give every figure/table with an id its label, so text can say {fig:id} before the figure."""
        f = t = 0
        for b in blocks:
            opts = b[-1] if isinstance(b[-1], dict) else {}
            if b[0] == "fig":
                f += 1
                if "id" in opts:
                    self.refs[f"fig:{opts['id']}"] = f"{self.chapter}.{f}"
            elif b[0] == "table":
                t += 1
                if "id" in opts:
                    self.refs[f"tab:{opts['id']}"] = f"{self.chapter}.{t}"

    def sub(self, text):
        return re.sub(r"\{((?:fig|tab):[a-z0-9_]+)\}", lambda m: self.refs[m.group(1)], text)

    def render(self, blocks):
        self.refs = {}
        self.number(blocks)
        for b in blocks:
            kind = b[0]
            if kind in ("p", "ul", "ol"):
                b = (kind, self.sub(b[1]) if kind == "p" else [self.sub(x) for x in b[1]])
            if kind == "p":
                self.para(b[1])
            elif kind == "h2":
                self.h2(b[1])
            elif kind == "h3":
                self.h3(b[1])
            elif kind == "ul":
                self.bullets(b[1])
            elif kind == "ol":
                self.bullets(b[1], numbered=True)
            elif kind == "fig":
                path, caption = b[1], b[2]
                opts = b[3] if len(b) > 3 else {}
                if "crop" in opts:
                    path = crop_image(REPO / path, opts["crop"], Path(path).stem + "_crop")
                elif "pair" in opts:
                    for dst, (src_, box_) in opts.get("crops", {}).items():
                        crop_image(REPO / src_, box_, Path(dst).stem)
                    path = side_by_side([REPO / x for x in opts["pair"]], opts["name"])
                else:
                    path = REPO / path
                self.figure(path, caption, opts.get("width", 6.0), opts.get("max_h", 7.6))
            elif kind == "table":
                self.table(b[1], b[2], b[3], b[4] if len(b) > 4 and not isinstance(b[4], dict) else None)
            elif kind == "break":
                self.page_break()
            else:
                raise ValueError(kind)

    # ---------- page-numbered lists
    def _leader_line(self, text, page, indent=0.0, bold=False, size=12):
        p = self.doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(indent); p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.line_spacing = 1.2
        p.paragraph_format.tab_stops.add_tab_stop(Inches(TEXT_W), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        r = p.add_run(f"{text}\t{page}"); r.bold = bold; r.font.size = Pt(size)

    def toc_block(self, entries):
        for lvl, text, key in entries:
            self._leader_line(text, self.pages.get(key, "00"), indent=[0, 0.3, 0.6][lvl], bold=(lvl == 0),
                              size=12 if lvl < 2 else 11)

    def fig_block(self, entries):
        for label, caption, key in entries:
            cap = caption if len(caption) < 92 else caption[:88].rsplit(" ", 1)[0] + "..."
            self._leader_line(f"{label}: {cap}", self.pages.get(key, "00"), size=11)


# ------------------------------------------------------------------ document assembly
def build(pages, toc_entries=None, lof_entries=None, lot_entries=None):
    R = Report(pages)
    d = R.doc

    # ---- cover page (format of the departmental template)
    cover = [("IT PROJECT REPORT ON", 14, 0), ("STUDENT INDUSTRIAL WORK EXPERIENCE SCHEME (SIWES)", 14, 0),
             ("UNDERTAKEN AT", 14, 18), ("ACADEVA LIMITED", 16, 30), ("BY", 14, 18),
             ("JOSEPH COSMOS ANIEBIET", 15, 0), ("22/SC/CO/1301", 14, 30),
             ("DEPARTMENT OF COMPUTER SCIENCE,", 13, 0), ("FACULTY OF COMPUTING,", 13, 0), ("UNIVERSITY OF UYO, UYO.", 13, 30),
             ("SUBMITTED TO", 14, 18),
             ("DEPARTMENT OF COMPUTER SCIENCE,", 13, 0), ("FACULTY OF COMPUTING,", 13, 0), ("UNIVERSITY OF UYO, UYO.", 13, 30),
             ("IN PARTIAL FULFILLMENT OF THE REQUIREMENT FOR THE AWARD OF A BACHELOR OF SCIENCE (B.Sc.) DEGREE IN COMPUTER SCIENCE.", 13, 36)]
    for text, size, after in cover:
        p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(after); p.paragraph_format.line_spacing = 1.3
        r = p.add_run(text); r.bold = True; r.font.size = Pt(size)
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run("OCTOBER, 2026."); r.bold = True; r.font.size = Pt(13)
    R.page_break()

    # ---- project title page
    for _ in range(8):
        d.add_paragraph()
    for text, size in (("IT PROJECT ON", 16), ("ARACARE HEALTH OS:", 18),
                       ("AN OFFLINE-FIRST HOSPITAL OPERATING SYSTEM WITH INTELLIGENT PATIENT QUEUING, "
                        "ELECTRONIC HEALTH RECORDS AND A LOCAL AI VOICE SCRIBE", 15)):
        p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(14)
        r = p.add_run(text); r.bold = True; r.font.size = Pt(size)
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(40)
    r = p.add_run("BY\nJOSEPH COSMOS ANIEBIET\n22/SC/CO/1301"); r.font.size = Pt(13); r.bold = True

    # ---- prelims (roman numerals continue from the cover = i, title = ii)
    sec = d.add_section(WD_SECTION.NEW_PAGE)
    sec.footer.is_linked_to_previous = False
    page_numbers(sec, "lowerRoman", 3)
    fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER; add_field(fp.add_run(), "PAGE")

    for i, (title, key, paras) in enumerate(PRELIMS):
        R.h_plain(title, key)
        if key == "dedication":
            for _ in range(4):
                d.add_paragraph()
            for t in paras:
                R.para(t, align="center", italic=True)
        else:
            for t in paras:
                if t.startswith("Keywords:"):
                    R.para(t, italic=True)
                else:
                    R.para(t)

    R.h_plain("Table of Contents", None)
    toc_at = len(d.paragraphs)
    R.h_plain("List of Figures", None)
    lof_at = len(d.paragraphs)
    R.h_plain("List of Tables", None)
    lot_at = len(d.paragraphs)

    # ---- body (arabic numbering restarts at 1)
    sec = d.add_section(WD_SECTION.NEW_PAGE)
    sec.footer.is_linked_to_previous = False
    page_numbers(sec, "decimal", 1)
    fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER; add_field(fp.add_run(), "PAGE")

    for i, (num, title, blocks) in enumerate(CHAPTERS):
        R.h1(num, title)
        R.render(blocks)

    R.chapter = 0
    R.h_plain("References", "references")
    for ref in REFERENCES:
        p = R.para(ref, align="left")
        p.paragraph_format.left_indent = Inches(0.5); p.paragraph_format.first_line_indent = Inches(-0.5)
    R.h_plain("Appendices", "appendices")
    R.chapter = "A"
    R.render(APPENDICES)

    # ---- fill the lists now that every entry is known (use pass-1 entries so line counts match)
    def insert_at(index, fill):
        anchor = d.paragraphs[index - 1]
        start = len(d.paragraphs)
        fill()
        new = d.paragraphs[start:]
        for p in reversed(new):
            anchor._p.addnext(p._p)
    # insert from the bottom up so earlier indices stay valid
    insert_at(lot_at, lambda: R.fig_block(lot_entries or R.lot))
    insert_at(lof_at, lambda: R.fig_block(lof_entries or R.lof))
    toc = [(0, "Dedication", "dedication"), (0, "Acknowledgement", "acknowledgement"), (0, "Abstract", "abstract"),
           (0, "List of Figures", "lof"), (0, "List of Tables", "lot")]
    insert_at(toc_at, lambda: R.toc_block(toc + [e for e in (toc_entries or R.toc) if e[2] not in
                                                 ("dedication", "acknowledgement", "abstract")]))
    return R


# ------------------------------------------------------------------ pagination (pass 1 -> page numbers)
def to_pdf(docx_path):
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(docx_path.parent), str(docx_path)],
                   check=True, capture_output=True, timeout=600)
    return docx_path.with_suffix(".pdf")


def roman(n):
    vals = [(10, "x"), (9, "ix"), (5, "v"), (4, "iv"), (1, "i")]
    out = ""
    for v, s in vals:
        while n >= v:
            out += s; n -= v
    return out


def norm(s):
    return re.sub(r"[^a-z0-9]", "", s.lower())


def locate(pdf_path, R):
    with pdfplumber.open(pdf_path) as pdf:
        texts = [norm(p.extract_text() or "") for p in pdf.pages]
        firsts = [norm((p.extract_text() or "").strip().split("\n")[0]) for p in pdf.pages]
    body0 = next(i for i, f in enumerate(firsts) if f.startswith("chapterone"))
    pages = {}

    def find(needle, start, end=None, first_line=False):
        n = norm(needle)
        for i in range(start, end or len(texts)):
            if (firsts[i].startswith(n) if first_line else n in texts[i]):
                return i
        return None

    for key, title in (("dedication", "dedication"), ("acknowledgement", "acknowledgement"), ("abstract", "abstract"),
                       ("lof", "list of figures"), ("lot", "list of tables")):
        i = find(title, 2, body0, first_line=True)
        if i is not None:
            pages[key] = roman(i + 1)
    for lvl, text, key in R.toc:
        if key.startswith("h1-"):
            i = find("chapter " + ["", "one", "two", "three"][int(key[3:])], body0, first_line=True)
        elif key in ("references", "appendices"):
            i = find(text, body0, first_line=True)
        else:
            i = find(text, body0)
        if i is not None:
            pages[key] = str(i - body0 + 1)
    for label, caption, key in R.lof + R.lot:
        i = find(f"{label}: {caption[:40]}", body0)
        if i is not None:
            pages[key] = str(i - body0 + 1)
    return pages, len(texts), body0


def main():
    R1 = build({})
    p1 = TMP / "pass1.docx"; R1.doc.save(p1)
    pages, n, body0 = locate(to_pdf(p1), R1)
    missing = [e[2] for e in R1.toc + R1.lof + R1.lot if e[2] not in pages]
    print(f"pass 1: {n} pages ({body0} preliminary), unresolved: {missing}")
    R2 = build(pages, R1.toc, R1.lof, R1.lot)
    final = OUT / f"{NAME}.docx"; R2.doc.save(final)
    pdf = to_pdf(final)
    pages2, n2, body2 = locate(pdf, R2)
    drift = {k: (pages[k], pages2.get(k)) for k in pages if pages2.get(k) != pages[k]}
    print(f"final: {n2} pages ({body2} preliminary + {n2 - body2} body). page drift: {drift or 'none'}")
    json.dump({"pages": n2, "prelim": body2, "figures": len(R2.lof), "tables": len(R2.lot)},
              open(OUT / "build_summary.json", "w"), indent=2)


if __name__ == "__main__":
    main()
